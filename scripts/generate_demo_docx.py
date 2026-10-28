import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def create_demo_docx(output_path="DEMONSTRATION_DOCUMENT.docx"):
    doc = docx.Document()

    # Set 0.75-inch margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(37, 99, 235)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 41, 59)

    def add_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(10)

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        if bold_prefix:
            r1 = p.add_run(bold_prefix + " ")
            r1.font.name = "Calibri"
            r1.font.size = Pt(9.5)
            r1.font.bold = True
        r2 = p.add_run(text)
        r2.font.name = "Calibri"
        r2.font.size = Pt(9.5)

    # Title & Header
    add_title("AI-Powered Portfolio Management & Dynamic Resume Engine")
    add_subtitle("Comprehensive System Demonstration & Architectural Walkthrough Document")
    add_p("Candidate & Systems Architect: Raj Bharat Rasal")
    add_p("System Version: 1.0.0 (FastAPI, SQLite 3, Google Gemini API, Google Sheets Apps Script)")
    add_p("Demonstration Video Assets: demo_video/portfolio_system_demonstration.mp4 (and .webm)")

    # Section 1
    add_h1("1. Executive Summary")
    add_p("This document provides an end-to-end architectural demonstration of the AI-Powered Portfolio Management & Dynamic Resume Generator. The platform maintains a single source of truth for all professional qualifications, systems projects, whitepapers, achievements, and competencies in a normalized SQLite relational database. It integrates the Google Gemini API to parse unstructured text, match candidate qualifications against arbitrary Job Descriptions (JDs), conduct gap analyses, and generate customized, ATS-compliant resumes in both PDF and DOCX formats without hallucinating information.")
    add_p("Additionally, the system synchronizes seamlessly with Google Sheets via a multi-tab workbook export (exported_portfolio.xlsx) and a live two-way Webhook powered by Google Apps Script.")

    # Section 2
    add_h1("2. Relational Database Architecture (DB Deep Dive)")
    add_p("The system is backed by a structured relational SQLite 3 database (portfolio.db) managed through SQLAlchemy ORM. The database enforces typed columns, foreign keys, and referential constraints across 9 specialized tables:")
    add_bullet("1. profiles:", "Stores candidate contact details (email, phone, location), portfolio links, GitHub, LinkedIn, and professional summary.")
    add_bullet("2. projects:", "Stores production and distributed systems projects: SATURDAY/SUNDAY (tensor sharding compute network), F.R.I.D.A.Y (decoupled multi-agent backend), EduOS (curiosity engine), and SecureVerse (zero-trust architecture).")
    add_bullet("3. skills:", "Contains 29 categorized skills across Languages, Distributed Systems & Architecture, AI/ML, Backend & Cloud, Security & Systems, and Soft Skills.")
    add_bullet("4. educations:", "Maintains educational history: B.Tech in CSE at ITM Skills University (2025-2029) and CBSE 10th & 12th records.")
    add_bullet("5. experiences:", "Records engineering practice and tutoring (taught 50+ students in Math & Physics).")
    add_bullet("6. publications:", "Maintains high-level architecture specifications: HAMIA (Hierarchical Adaptive Multi-Model Intelligence Architecture) and TWELVE (AI Assessment System).")
    add_bullet("7. achievements:", "Captures competition awards: Deal Dome 3rd Place for SecureVerse, HackerOne OPPO Find X3 Pro bug bounty disclosure.")
    add_bullet("8. certifications:", "Stores credentials in distributed systems, security research, and multi-agent pipelines.")
    add_bullet("9. languages:", "Tracks language proficiencies: English, Hindi, Marathi.")

    # Section 3
    add_h1("3. Google Sheets Integration Deep Dive")
    add_p("The system treats Google Sheets as a premier, real-time accessible mirror of the portfolio database:")
    add_bullet("One-Click Multi-Tab XLSX Export:", "Generates an Excel workbook (sample_data/exported_portfolio.xlsx) with 9 separate sheets matching the database tables. Features styled slate headers, alternating row fills, thin grid borders, and auto-adjusted column dimensions. Can be directly imported into Google Sheets via File > Import.")
    add_bullet("Live 2-Way Sync via Google Apps Script:", "Users deploy the provided Google Apps Script (google_apps_script/Code.gs) in their Google Sheet. Pushing the database from the web application transmits the JSON payload via Webhook, instantly writing and formatting all 9 sheets in Google Sheets cloud.")
    add_bullet("Interactive In-App Spreadsheet Preview:", "The Google Sheets tab inside the application provides a live spreadsheet viewer with sheet tab switchers, column headers (A, B, C...), and numbered rows (1, 2, 3...).")

    # Section 4
    add_h1("4. Gemini API Integration & Ground-Truth Verification Invariant")
    add_p("The platform incorporates the Google Gemini API (gemini-2.0-flash / gemini-1.5-flash) for AI intelligence:")
    add_bullet("Unstructured Text & Resume Parser:", "Extracts structured entity records from raw text and maps them to database schema fields with 1-click import.")
    add_bullet("Job Description Keyword & Gap Analysis:", "Evaluates candidate skills against JD requirements, computes an objective Match Score (0–100%), extracts domain keywords, and identifies missing vs. matched skills.")
    add_bullet("Strict Ground-Truth Invariant:", "The system enforces an inviolable rule preventing hallucination. The AI is bounded strictly to verified candidate records in portfolio.db and will never invent fictitious degrees, companies, or credentials.")

    # Section 5
    add_h1("5. Dynamic Resume Generator & Dual Customization")
    add_p("The platform demonstrates dynamic resume generation by evaluating two distinct industry roles from the exact same portfolio:")
    add_bullet("Job 1: Backend & Distributed Systems Architect:", "Generates a resume highlighting SATURDAY/SUNDAY tensor sharding, C++, Python, zero-trust protocols, and the HAMIA whitepaper. (88% Match Score).")
    add_bullet("Job 2: AI Systems & Multi-Agent Infrastructure Engineer:", "Generates a resume highlighting F.R.I.D.A.Y multi-agent mediator backend, EduOS curiosity engine, intent routing, and prompt chaining. (86% Match Score).")
    add_bullet("Dual Format Download:", "Every customized resume is downloadable as an ATS-compliant PDF (ReportLab) or editable Word document (.docx).")

    # Section 6
    add_h1("6. Demonstration Video Asset Reference")
    add_p("An automated live walkthrough video has been recorded in high-definition (1280x720, 25 FPS) and is available at:")
    add_bullet("MP4 File:", "demo_video/portfolio_system_demonstration.mp4")
    add_bullet("WebM File:", "demo_video/portfolio_system_demonstration.webm")

    doc.save(output_path)
    print(f"Generated {output_path} successfully!")

if __name__ == "__main__":
    create_demo_docx()

