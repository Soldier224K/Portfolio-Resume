import io
from typing import Dict, Any
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def add_horizontal_rule(paragraph, color_hex="CCCCCC"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)

def generate_docx_resume(resume_data: Dict[str, Any], output_path: str = None) -> io.BytesIO:
    """
    Generate an ATS-compliant, cleanly formatted DOCX Word document.
    """
    doc = docx.Document()

    # Set 0.5 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    profile = resume_data.get("profile", {})
    target_title = resume_data.get("target_title", profile.get("title", ""))

    # 1. Header Name
    p_name = doc.add_paragraph()
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(2)
    run_name = p_name.add_run(profile.get("name", "Raj Bharat Rasal").upper())
    run_name.font.name = "Calibri"
    run_name.font.size = Pt(20)
    run_name.font.bold = True
    run_name.font.color.rgb = RGBColor(15, 23, 42)

    # Target Role Title
    if target_title:
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(0)
        p_title.paragraph_format.space_after = Pt(2)
        run_title = p_title.add_run(target_title)
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(11)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(37, 99, 235)

    # Contact line
    contact_parts = []
    if profile.get("location"):
        contact_parts.append(profile["location"])
    if profile.get("phone"):
        contact_parts.append(profile["phone"])
    if profile.get("email"):
        contact_parts.append(profile["email"])
    if profile.get("linkedin"):
        contact_parts.append(profile["linkedin"].replace("https://", ""))
    if profile.get("github"):
        contact_parts.append(profile["github"].replace("https://", ""))

    p_contact = doc.add_paragraph()
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.space_after = Pt(6)
    run_contact = p_contact.add_run("  |  ".join(contact_parts))
    run_contact.font.name = "Calibri"
    run_contact.font.size = Pt(9)
    run_contact.font.color.rgb = RGBColor(71, 85, 105)
    add_horizontal_rule(p_contact, "94A3B8")

    def add_section_header(title):
        p_h = doc.add_paragraph()
        p_h.paragraph_format.space_before = Pt(8)
        p_h.paragraph_format.space_after = Pt(3)
        run_h = p_h.add_run(title.upper())
        run_h.font.name = "Calibri"
        run_h.font.size = Pt(10.5)
        run_h.font.bold = True
        run_h.font.color.rgb = RGBColor(15, 23, 42)
        add_horizontal_rule(p_h, "CBD5E1")

    # 2. Professional Summary
    summary = resume_data.get("custom_summary") or profile.get("summary")
    if summary:
        add_section_header("Professional Summary")
        p_sum = doc.add_paragraph()
        p_sum.paragraph_format.space_before = Pt(2)
        p_sum.paragraph_format.space_after = Pt(4)
        run_sum = p_sum.add_run(summary)
        run_sum.font.name = "Calibri"
        run_sum.font.size = Pt(9.5)

    # 3. Technical Skills
    skills = resume_data.get("skills", [])
    if skills:
        add_section_header("Technical & Architectural Skills")
        categories = {}
        for s in skills:
            cat = s.get("category", "General")
            categories.setdefault(cat, []).append(s["name"])

        for cat, s_names in categories.items():
            p_s = doc.add_paragraph()
            p_s.paragraph_format.space_before = Pt(1)
            p_s.paragraph_format.space_after = Pt(2)
            run_c = p_s.add_run(f"{cat}: ")
            run_c.font.name = "Calibri"
            run_c.font.size = Pt(9.5)
            run_c.font.bold = True
            run_val = p_s.add_run(", ".join(s_names))
            run_val.font.name = "Calibri"
            run_val.font.size = Pt(9.5)

    # 4. Featured Projects
    projects = resume_data.get("projects", [])
    if projects:
        add_section_header("Featured Systems & Projects")
        for p in projects:
            p_item = doc.add_paragraph()
            p_item.paragraph_format.space_before = Pt(4)
            p_item.paragraph_format.space_after = Pt(1)
            
            run_t = p_item.add_run(f"{p.get('title')} — {p.get('subtitle_role', '')}")
            run_t.font.name = "Calibri"
            run_t.font.size = Pt(9.5)
            run_t.font.bold = True
            
            if p.get("start_date"):
                date_str = f"   ({p.get('start_date')} - {p.get('end_date', 'Present')})"
                run_d = p_item.add_run(date_str)
                run_d.font.name = "Calibri"
                run_d.font.size = Pt(8.5)
                run_d.font.italic = True
                run_d.font.color.rgb = RGBColor(100, 116, 139)

            if p.get("tech_stack"):
                p_ts = doc.add_paragraph()
                p_ts.paragraph_format.space_before = Pt(0)
                p_ts.paragraph_format.space_after = Pt(2)
                run_ts = p_ts.add_run(f"Tech Stack: {p.get('tech_stack')}")
                run_ts.font.name = "Calibri"
                run_ts.font.size = Pt(8.5)
                run_ts.font.italic = True
                run_ts.font.color.rgb = RGBColor(71, 85, 105)

            bullets = [b.strip() for b in p.get("description", "").split("\n") if b.strip()]
            for b in bullets:
                clean_b = b.lstrip("•-* ").strip()
                bp = doc.add_paragraph(style='List Bullet')
                bp.paragraph_format.space_before = Pt(0)
                bp.paragraph_format.space_after = Pt(1.5)
                r_b = bp.add_run(clean_b)
                r_b.font.name = "Calibri"
                r_b.font.size = Pt(9)

            if p.get("architecture_highlights"):
                bp_arch = doc.add_paragraph(style='List Bullet')
                bp_arch.paragraph_format.space_before = Pt(0)
                bp_arch.paragraph_format.space_after = Pt(2)
                r_tag = bp_arch.add_run("Key Architecture: ")
                r_tag.font.name = "Calibri"
                r_tag.font.size = Pt(9)
                r_tag.font.bold = True
                r_arch = bp_arch.add_run(p.get("architecture_highlights"))
                r_arch.font.name = "Calibri"
                r_arch.font.size = Pt(9)

    # 5. Experience
    experiences = resume_data.get("experiences", [])
    if experiences:
        add_section_header("Experience & Leadership")
        for exp in experiences:
            p_exp = doc.add_paragraph()
            p_exp.paragraph_format.space_before = Pt(4)
            p_exp.paragraph_format.space_after = Pt(1)
            
            run_e = p_exp.add_run(f"{exp.get('company')} — {exp.get('role')}")
            run_e.font.name = "Calibri"
            run_e.font.size = Pt(9.5)
            run_e.font.bold = True

            date_str = f"   ({exp.get('start_date')} - {exp.get('end_date', 'Present')})"
            run_ed = p_exp.add_run(date_str)
            run_ed.font.name = "Calibri"
            run_ed.font.size = Pt(8.5)
            run_ed.font.italic = True
            run_ed.font.color.rgb = RGBColor(100, 116, 139)

            bullets = [b.strip() for b in exp.get("description", "").split("\n") if b.strip()]
            for b in bullets:
                clean_b = b.lstrip("•-* ").strip()
                bp = doc.add_paragraph(style='List Bullet')
                bp.paragraph_format.space_before = Pt(0)
                bp.paragraph_format.space_after = Pt(1.5)
                r_b = bp.add_run(clean_b)
                r_b.font.name = "Calibri"
                r_b.font.size = Pt(9)

    # 6. Publications / Whitepapers
    publications = resume_data.get("publications", [])
    if publications:
        add_section_header("System Design Specifications & Whitepapers")
        for pub in publications:
            p_pub = doc.add_paragraph(style='List Bullet')
            p_pub.paragraph_format.space_before = Pt(1)
            p_pub.paragraph_format.space_after = Pt(2)
            r_pt = p_pub.add_run(f"{pub.get('title')}: ")
            r_pt.font.name = "Calibri"
            r_pt.font.size = Pt(9)
            r_pt.font.bold = True
            r_pa = p_pub.add_run(pub.get('abstract', ''))
            r_pa.font.name = "Calibri"
            r_pa.font.size = Pt(9)

    # 7. Education
    education = resume_data.get("education", [])
    if education:
        add_section_header("Education")
        for edu in education:
            p_edu = doc.add_paragraph()
            p_edu.paragraph_format.space_before = Pt(3)
            p_edu.paragraph_format.space_after = Pt(1)
            r_inst = p_edu.add_run(f"{edu.get('institution')} — {edu.get('degree')}")
            r_inst.font.name = "Calibri"
            r_inst.font.size = Pt(9.5)
            r_inst.font.bold = True

            d_str = f"   ({edu.get('start_year')} - {edu.get('end_year')})"
            r_ed = p_edu.add_run(d_str)
            r_ed.font.name = "Calibri"
            r_ed.font.size = Pt(8.5)
            r_ed.font.italic = True

            if edu.get("highlights"):
                bp_eh = doc.add_paragraph(style='List Bullet')
                bp_eh.paragraph_format.space_before = Pt(0)
                bp_eh.paragraph_format.space_after = Pt(1.5)
                r_eh = bp_eh.add_run(edu.get("highlights"))
                r_eh.font.name = "Calibri"
                r_eh.font.size = Pt(9)

    # 8. Achievements
    achievements = resume_data.get("achievements", [])
    if achievements:
        add_section_header("Achievements & Security Research")
        for a in achievements:
            p_a = doc.add_paragraph(style='List Bullet')
            p_a.paragraph_format.space_before = Pt(0)
            p_a.paragraph_format.space_after = Pt(1.5)
            r_at = p_a.add_run(f"{a.get('title')}: ")
            r_at.font.name = "Calibri"
            r_at.font.size = Pt(9)
            r_at.font.bold = True
            r_ad = p_a.add_run(a.get('description', ''))
            r_ad.font.name = "Calibri"
            r_ad.font.size = Pt(9)

    # 9. Certifications
    certifications = resume_data.get("certifications", [])
    if certifications:
        add_section_header("Certifications & Accreditations")
        for c in certifications:
            p_c = doc.add_paragraph(style='List Bullet')
            p_c.paragraph_format.space_before = Pt(0)
            p_c.paragraph_format.space_after = Pt(1.5)
            r_ct = p_c.add_run(f"{c.get('name')} ")
            r_ct.font.name = "Calibri"
            r_ct.font.size = Pt(9)
            r_ct.font.bold = True
            r_co = p_c.add_run(f"— {c.get('issuing_org')} ({c.get('issue_date')})")
            r_co.font.name = "Calibri"
            r_co.font.size = Pt(9)

    buffer = io.BytesIO()
    if output_path:
        doc.save(output_path)
    doc.save(buffer)
    buffer.seek(0)
    return buffer

