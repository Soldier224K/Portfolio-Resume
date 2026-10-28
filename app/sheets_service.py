import os
import io
import json
import httpx
from typing import Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Default path for exported portfolio spreadsheet
EXPORT_DIR = os.path.join(os.path.dirname(__file__), "..", "sample_data")
EXPORT_XLSX_PATH = os.path.join(EXPORT_DIR, "exported_portfolio.xlsx")

HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")  # Slate 800
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ROW_FONT = Font(name="Calibri", size=10, color="0F172A")
ALT_ROW_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")  # Slate 50
BORDER_THIN = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

def create_portfolio_workbook(portfolio_data: Dict[str, Any]) -> openpyxl.Workbook:
    """
    Generate a beautifully formatted, multi-tab Excel Workbook containing all
    categories of the user's structured portfolio database.
    Directly compatible with Google Sheets import.
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    def style_table(ws, headers, rows):
        ws.views.sheetView[0].showGridLines = True
        ws.append(headers)

        # Style Header
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER_THIN

        # Style Rows
        for row_idx, row_data in enumerate(rows, start=2):
            ws.append(row_data)
            is_alt = (row_idx % 2 == 1)
            for col_num in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_num)
                cell.font = ROW_FONT
                cell.border = BORDER_THIN
                if is_alt:
                    cell.fill = ALT_ROW_FILL
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        # Set Column Widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                first_line = val.split("\n")[0]
                max_len = max(max_len, min(len(first_line), 50))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # 1. Profile Sheet
    ws_prof = wb.create_sheet(title="Profile")
    p = portfolio_data.get("profile", {}) or {}
    prof_headers = ["Field", "Value"]
    prof_rows = [
        ["Name", p.get("name", "")],
        ["Title", p.get("title", "")],
        ["Email", p.get("email", "")],
        ["Phone", p.get("phone", "")],
        ["Location", p.get("location", "")],
        ["LinkedIn", p.get("linkedin", "")],
        ["GitHub", p.get("github", "")],
        ["Portfolio URL", p.get("portfolio_url", "")],
        ["Professional Summary", p.get("summary", "")],
        ["Last Updated", p.get("updated_at", "")]
    ]
    style_table(ws_prof, prof_headers, prof_rows)

    # 2. Education Sheet
    ws_edu = wb.create_sheet(title="Education")
    edu_headers = ["ID", "Institution", "Degree", "Field of Study", "Start Year", "End Year", "Score", "Highlights"]
    edu_rows = [
        [
            e.get("id", ""),
            e.get("institution", ""),
            e.get("degree", ""),
            e.get("field_of_study", ""),
            e.get("start_year", ""),
            e.get("end_year", ""),
            e.get("score", ""),
            e.get("highlights", "")
        ]
        for e in portfolio_data.get("education", [])
    ]
    style_table(ws_edu, edu_headers, edu_rows)

    # 3. Skills Sheet
    ws_skills = wb.create_sheet(title="Skills")
    skills_headers = ["ID", "Skill Name", "Category", "Proficiency", "Highlighted"]
    skills_rows = [
        [
            s.get("id", ""),
            s.get("name", ""),
            s.get("category", ""),
            s.get("proficiency", ""),
            "Yes" if s.get("is_highlighted") else "No"
        ]
        for s in portfolio_data.get("skills", [])
    ]
    style_table(ws_skills, skills_headers, skills_rows)

    # 4. Experience Sheet
    ws_exp = wb.create_sheet(title="Experience")
    exp_headers = ["ID", "Company", "Role", "Location", "Start Date", "End Date", "Current", "Description", "Technologies"]
    exp_rows = [
        [
            ex.get("id", ""),
            ex.get("company", ""),
            ex.get("role", ""),
            ex.get("location", ""),
            ex.get("start_date", ""),
            ex.get("end_date", ""),
            "Yes" if ex.get("is_current") else "No",
            ex.get("description", ""),
            ex.get("technologies", "")
        ]
        for ex in portfolio_data.get("experiences", [])
    ]
    style_table(ws_exp, exp_headers, exp_rows)

    # 5. Projects Sheet
    ws_proj = wb.create_sheet(title="Projects")
    proj_headers = ["ID", "Title", "Role / Subtitle", "Tech Stack", "Description", "Architecture Highlights", "Repo Link", "Demo Link", "Featured"]
    proj_rows = [
        [
            pr.get("id", ""),
            pr.get("title", ""),
            pr.get("subtitle_role", ""),
            pr.get("tech_stack", ""),
            pr.get("description", ""),
            pr.get("architecture_highlights", ""),
            pr.get("repo_link", ""),
            pr.get("demo_link", ""),
            "Yes" if pr.get("is_featured") else "No"
        ]
        for pr in portfolio_data.get("projects", [])
    ]
    style_table(ws_proj, proj_headers, proj_rows)

    # 6. Certifications Sheet
    ws_cert = wb.create_sheet(title="Certifications")
    cert_headers = ["ID", "Certification Name", "Issuing Organization", "Issue Date", "Credential ID", "Credential URL"]
    cert_rows = [
        [
            c.get("id", ""),
            c.get("name", ""),
            c.get("issuing_org", ""),
            c.get("issue_date", ""),
            c.get("credential_id", ""),
            c.get("credential_url", "")
        ]
        for c in portfolio_data.get("certifications", [])
    ]
    style_table(ws_cert, cert_headers, cert_rows)

    # 7. Achievements Sheet
    ws_ach = wb.create_sheet(title="Achievements")
    ach_headers = ["ID", "Title", "Issuer", "Date", "Description", "Category"]
    ach_rows = [
        [
            a.get("id", ""),
            a.get("title", ""),
            a.get("issuer", ""),
            a.get("date", ""),
            a.get("description", ""),
            a.get("category", "")
        ]
        for a in portfolio_data.get("achievements", [])
    ]
    style_table(ws_ach, ach_headers, ach_rows)

    # 8. Publications Sheet
    ws_pub = wb.create_sheet(title="Publications")
    pub_headers = ["ID", "Title", "Type", "Publisher", "Publication Date", "URL", "Abstract"]
    pub_rows = [
        [
            pb.get("id", ""),
            pb.get("title", ""),
            pb.get("venue_type", ""),
            pb.get("publisher", ""),
            pb.get("publication_date", ""),
            pb.get("url", ""),
            pb.get("abstract", "")
        ]
        for pb in portfolio_data.get("publications", [])
    ]
    style_table(ws_pub, pub_headers, pub_rows)

    # 9. Languages Sheet
    ws_lang = wb.create_sheet(title="Languages")
    lang_headers = ["ID", "Language", "Proficiency"]
    lang_rows = [
        [l.get("id", ""), l.get("language", ""), l.get("proficiency", "")]
        for l in portfolio_data.get("languages", [])
    ]
    style_table(ws_lang, lang_headers, lang_rows)

    return wb

def export_portfolio_to_file(portfolio_data: Dict[str, Any], filepath: Optional[str] = None) -> str:
    """Generate and write the multi-tab XLSX file to disk."""
    path = filepath or EXPORT_XLSX_PATH
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb = create_portfolio_workbook(portfolio_data)
    wb.save(path)
    return os.path.abspath(path)

def export_portfolio_to_bytes(portfolio_data: Dict[str, Any]) -> io.BytesIO:
    """Generate the multi-tab XLSX file in an in-memory byte buffer."""
    wb = create_portfolio_workbook(portfolio_data)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

async def sync_to_google_sheets_webhook(webhook_url: str, portfolio_data: Dict[str, Any]) -> Dict[str, Any]:
    """
    Push structured portfolio records directly to a Google Apps Script Web App URL.
    The Google Apps Script automatically creates/updates all sheets with fresh records.
    """
    if not webhook_url or not webhook_url.startswith("http"):
        return {"success": False, "message": "Invalid Google Apps Script Webhook URL"}

    payload = {
        "action": "sync_portfolio",
        "timestamp": portfolio_data.get("profile", {}).get("updated_at"),
        "data": portfolio_data
    }

    try:
        async with httpx.AsyncClient(timeout=20.0, follow_redirects=True) as client:
            resp = await client.post(webhook_url, json=payload)
            if resp.status_code in [200, 201, 302]:
                return {
                    "success": True,
                    "status_code": resp.status_code,
                    "message": "Successfully synchronized portfolio database to Google Sheets!",
                    "response": resp.text[:300]
                }
            else:
                return {
                    "success": False,
                    "status_code": resp.status_code,
                    "message": f"Google Sheets webhook returned status {resp.status_code}: {resp.text[:200]}"
                }
    except Exception as e:
        return {"success": False, "message": f"Connection error syncing to Google Sheets: {str(e)}"}

